"""Recommandations par pays (CSR) du Semestre europeen : base ECFIN (export Qlik) -> une fiche Markdown par CSR.

Trois niveaux : document (un par Etat membre et par annee, `csrdoc-<ISO>-<annee>`, derive des recommandations),
recommandation (`csr-<ISO>-<annee>-<n>`) et sous-partie (`csr-<ISO>-<annee>-<n>.<m>`, base ECFIN seulement).

L'export a une ligne par CSR ("1") ou sous-partie ("1.2") et par version annuelle d'evaluation ("2025 - Semester
Cycle (June)") : texte et domaines de la derniere version, historique des evaluations garde. Les domaines CSR
(taxonomie ECFIN) restent dans le front matter et le texte : ce ne sont pas les domaines des projets.
"""

import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from tsi_agent.countries import iso_from_name, name_from_iso
from tsi_agent.render import render

SOURCE_URL = "https://ec.europa.eu/economy_finance/country-specific-recommendations-database/"
YEAR_FROM, YEAR_TO = 2016, 2025  # SRSP (2017) et TSI (2021) : CSR des annees precedant les demandes
COLUMNS = ("Country", "Year", "CSR Nbr", "Version", "Text", "Policy Area(s)", "EAR(s)", "MIP",
           "Annual/Multiannual", "Assessment")
_VERSION_YEAR = re.compile(r"^\s*(\d{4})")
_EMPTY = {"", "-"}
NOT_ASSESSED = "Not Assessed"
COUNCIL_STAGE = "Council recommendation"  # CSR adoptees (base ECFIN)
COMMISSION_STAGE = "Commission recommendation"  # propositions de la Commission, pas encore adoptees ni evaluees
SUMMARY_CHARS = 300


def _clean(value) -> str:
    text = " ".join(str(value if value is not None else "").split())
    return "" if text in _EMPTY else text


def _areas(value) -> list[str]:
    return [a for a in (_clean(p) for p in str(value or "").split(";")) if a]


def _version_year(version: str) -> int:
    match = _VERSION_YEAR.match(version)
    return int(match.group(1)) if match else 0


def _number_key(number: str) -> tuple[int, ...]:
    return tuple(int(n) if n.isdigit() else 0 for n in number.split("."))


def _latest(assessments: list[tuple[int, str]]) -> tuple[int, str] | None:
    """Derniere evaluation reelle : "Not Assessed" ne masque pas une evaluation anterieure."""
    real = [e for e in assessments if e[1].lower() != NOT_ASSESSED.lower()]
    return real[-1] if real else (assessments[-1] if assessments else None)


@dataclass
class Part:
    """Une CSR complete ou une sous-partie, a sa derniere version."""
    part_id: str
    text: str
    policy_areas: list[str]
    assessments: list[tuple[int, str]]  # (annee de la version, evaluation), chronologique

    @property
    def latest_assessment(self) -> tuple[int, str] | None:
        return _latest(self.assessments)

    @property
    def assessment(self) -> str | None:
        latest = self.latest_assessment
        return latest[1] if latest else None


@dataclass
class Csr:
    country: str
    year: int
    number: str
    full: Part | None
    subparts: list[Part] = field(default_factory=list)
    stage: str = COUNCIL_STAGE
    source_url: str = SOURCE_URL

    @property
    def csr_id(self) -> str:
        return f"csr-{self.country}-{self.year}-{self.number}"

    @property
    def doc_id(self) -> str:
        return document_id(self.country, self.year)

    @property
    def title(self) -> str:
        return f"{name_from_iso(self.country)} {self.year} CSR {self.number}"

    @property
    def text(self) -> str:
        return self.full.text if self.full else " ".join(p.text for p in self.subparts if p.text)

    @property
    def policy_areas(self) -> list[str]:
        full = self.full.policy_areas if self.full else []
        return list(dict.fromkeys([*full, *(a for p in self.subparts for a in p.policy_areas)]))

    @property
    def assessments(self) -> list[tuple[int, str]]:
        return self.full.assessments if self.full else []

    @property
    def latest_assessment(self) -> tuple[int, str] | None:
        return self.full.latest_assessment if self.full else None

    @property
    def assessment(self) -> str | None:
        return self.full.assessment if self.full else None


def document_id(country: str, year: int) -> str:
    return f"csrdoc-{country}-{year}"


@dataclass
class CsrDocument:
    """Document de recommandations d'un Etat membre pour une annee (27 par an)."""
    country: str
    year: int
    stage: str
    document_url: str
    recommendations: list[Csr]

    @property
    def doc_id(self) -> str:
        return document_id(self.country, self.year)

    @property
    def title(self) -> str:
        return f"{name_from_iso(self.country)} — {self.year} country-specific recommendations"

    @property
    def n_recommendations(self) -> int:
        return len(self.recommendations)


def build_documents(csrs: list[Csr]) -> list[CsrDocument]:
    """Documents par (pays, annee), recommandations par numero ; etape et lien de la premiere recommandation."""
    grouped: dict[tuple[str, int], list[Csr]] = defaultdict(list)
    for csr in csrs:
        grouped[(csr.country, csr.year)].append(csr)
    docs = []
    for (country, year), recs in sorted(grouped.items()):
        recs = sorted(recs, key=lambda c: _number_key(c.number))
        docs.append(CsrDocument(country, year, recs[0].stage, recs[0].source_url, recs))
    return docs


def read_rows(path: Path) -> list[dict]:
    """Lignes de l'export xlsx (premiere ligne = en-tetes)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        rows = wb.active.iter_rows(values_only=True)
        header = [str(h).strip() for h in next(rows)]
        return [dict(zip(header, r)) for r in rows if any(v is not None for v in r)]
    finally:
        wb.close()


def _part(part_id: str, versions: list[dict]) -> Part:
    versions = sorted(versions, key=lambda r: _version_year(_clean(r.get("Version"))))
    last = versions[-1]
    assessments = [(_version_year(_clean(r.get("Version"))), _clean(r.get("Assessment"))) for r in versions]
    return Part(part_id, _clean(last.get("Text")), _areas(last.get("Policy Area(s)")),
                [(y, a) for y, a in assessments if a])


def build_csrs(rows: list[dict], year_from: int = YEAR_FROM, year_to: int = YEAR_TO) -> list[Csr]:
    """CSR des Etats membres (pays inconnus, dont le Royaume-Uni, ignores) adoptees entre year_from et year_to."""
    grouped: dict[tuple[str, int, str], dict[str, list[dict]]] = defaultdict(lambda: defaultdict(list))
    for r in rows:
        iso = iso_from_name(_clean(r.get("Country")))
        number = _clean(r.get("CSR Nbr"))
        try:
            year = int(float(_clean(r.get("Year"))))
        except ValueError:
            continue
        if not iso or not re.fullmatch(r"\d+(\.\d+)?", number) or not year_from <= year <= year_to:
            continue
        main, _, sub = number.partition(".")
        grouped[(iso, year, main)][number if sub and sub != "0" else main].append(r)
    csrs = []
    for (iso, year, main), parts in sorted(grouped.items()):
        prefix = f"csr-{iso}-{year}-"
        full = _part(prefix + main, parts[main]) if main in parts else None
        subs = [_part(prefix + n, parts[n]) for n in sorted((n for n in parts if n != main), key=_number_key)]
        csrs.append(Csr(iso, year, main, full, subs))
    return csrs


def _subpart_line(p: Part) -> str:
    facts = []
    if p.policy_areas:
        facts.append(f"Policy areas: {'; '.join(p.policy_areas)}")
    if p.latest_assessment:
        year, value = p.latest_assessment
        facts.append(f"Assessment ({year}): {value}")
    return f"- {p.part_id}: {p.text}" + (f" ({' · '.join(facts)})" if facts else "")


def render_csr(csr: Csr) -> str:
    front = {"csr_id": csr.csr_id, "title": csr.title, "countries": [csr.country], "year": csr.year,
             "number": csr.number, "policy_areas": csr.policy_areas, "assessment": csr.assessment,
             "subparts": [p.part_id for p in csr.subparts], "source_url": csr.source_url,
             "doc_id": csr.doc_id, "stage": csr.stage, "csr_level": "recommendation"}
    lines = [f"# {csr.title}", "", csr.text]
    if csr.stage == COMMISSION_STAGE:
        lines += ["", f"{COMMISSION_STAGE} of {csr.year} (not yet adopted by the Council). Not yet assessed. "
                      f"Source: {csr.source_url}"]
    if csr.policy_areas:
        lines += ["", f"Policy areas: {'; '.join(csr.policy_areas)}."]
    if csr.assessments:
        last_year, last = csr.latest_assessment
        # dans l'introduction aussi : search_csrs ne garde souvent que ce morceau par CSR
        lines += ["", f"Latest assessment ({last_year}): {last}."]
        history = "; ".join(f"{y}: {a}" for y, a in csr.assessments)
        lines += ["", "## Assessment", "", f"Latest assessment ({last_year}): {last}. History: {history}."]
    if csr.subparts:
        lines += ["", "## Sub-parts", ""] + [_subpart_line(p) for p in csr.subparts]
    return render(front, "\n".join(lines))


def _summary(text: str, limit: int = SUMMARY_CHARS) -> str:
    text = " ".join(text.split())
    if len(text) <= limit:
        return text
    cut = text[:limit].rsplit(" ", 1)[0].rstrip(",;:")
    return cut + "…"


def _recommendation_line(csr: Csr) -> str:
    if csr.latest_assessment:
        year, value = csr.latest_assessment
        facts = [f"Latest assessment ({year}): {value}"]
    else:
        facts = ["Not yet assessed" if csr.stage == COMMISSION_STAGE else "No assessment"]
    if csr.subparts:
        facts.append(f"{len(csr.subparts)} sub-parts")
    return f"- {csr.csr_id}: {_summary(csr.text)} ({'; '.join(facts)})"


def render_document(doc: CsrDocument) -> str:
    front = {"doc_id": doc.doc_id, "title": doc.title, "countries": [doc.country], "year": doc.year,
             "stage": doc.stage, "n_recommendations": doc.n_recommendations,
             "recommendations": [c.csr_id for c in doc.recommendations], "source_url": doc.document_url,
             "csr_level": "document"}
    if doc.stage == COMMISSION_STAGE:
        intro = (f"{COMMISSION_STAGE} for a Council recommendation on the economic, social, employment, structural "
                 f"and budgetary policies of {name_from_iso(doc.country)} ({doc.year} European Semester): "
                 f"{doc.n_recommendations} recommendations, not yet adopted by the Council nor assessed. "
                 f"Document: {doc.document_url}")
    else:
        intro = (f"{COUNCIL_STAGE}s addressed to {name_from_iso(doc.country)} in the {doc.year} European Semester: "
                 f"{doc.n_recommendations} recommendations. Texts and assessments from the Commission's CSR "
                 f"database: {doc.document_url}")
    lines = [f"# {doc.title}", "", intro, "", "## Recommendations", ""]
    lines += [_recommendation_line(c) for c in doc.recommendations]
    return render(front, "\n".join(lines))


def write_documents(docs: list[CsrDocument], out_dir: Path) -> int:
    """Une fiche par document dans out_dir (fiches csrdoc-*.md precedentes supprimees, autres fichiers gardes)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("csrdoc-*.md"):
        old.unlink()
    for doc in docs:
        (out_dir / f"{doc.doc_id}.md").write_text(render_document(doc), encoding="utf-8", newline="\n")
    return len(docs)


def write_csrs(csrs: list[Csr], out_dir: Path) -> int:
    """Une fiche par CSR dans out_dir (fiches csr-*.md precedentes supprimees, autres fichiers gardes)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("csr-*.md"):
        old.unlink()
    for csr in csrs:
        (out_dir / f"{csr.csr_id}.md").write_text(render_csr(csr), encoding="utf-8", newline="\n")
    return len(csrs)
